#!/usr/bin/env python3
"""Create an edit-stable DBR DOCX with unified Korean/Latin fonts.

The source DOCX is treated as an immutable layout master.  This postprocessor
updates WordprocessingML directly so it does not rebuild tables, images,
sections, or pagination:

* all Korean/East Asian text -> NanumMyeongjo
* all Latin, digits, and complex-script text -> Times New Roman
* document defaults, styles, numbering, headers/footers, and direct run
  formatting receive the same mapping
* NanumMyeongjo Regular/Bold are embedded with editable-embedding permission
* East Asian document-grid snapping is disabled to avoid editor-dependent
  line jumps
* paragraph/section property order is normalized to the WordprocessingML
  schema, and two-column widths are written explicitly for Word compatibility

The DBR Rule 7 point sizes are intentionally preserved from the source DOCX.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import uuid
import zipfile
from copy import deepcopy
from pathlib import Path

from fontTools.ttLib import TTFont
from lxml import etree


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT = ROOT / "2026_KIISE" / "manuscript"
SOURCE = MANUSCRIPT / "kiise_dbr_manuscript_v6_submission_revision.docx"
OUTPUT = MANUSCRIPT / "kiise_dbr_manuscript_v6_submission_revision_editable.docx"
REPORT = (
    ROOT
    / "2026_KIISE"
    / "paper_assets"
    / "20260718_editable_docx"
    / "font_unification_report.json"
)

KOREAN_FONT = "NanumMyeongjo"
LATIN_FONT = "Times New Roman"
KOREAN_REGULAR = Path("/usr/share/fonts/truetype/nanum/NanumMyeongjo.ttf")
KOREAN_BOLD = Path("/usr/share/fonts/truetype/nanum/NanumMyeongjoBold.ttf")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
W = f"{{{W_NS}}}"
R = f"{{{R_NS}}}"
A = f"{{{A_NS}}}"
REL = f"{{{REL_NS}}}"
CT = f"{{{CT_NS}}}"
NS = {"w": W_NS, "r": R_NS, "a": A_NS}

XML_PART_RE = re.compile(
    r"^word/(?:document|styles|numbering|footnotes|endnotes|comments|"
    r"header\d+|footer\d+)\.xml$"
)

STALE_ANSWER = (
    "VRU 600문항의 고정 생성기에서 closed-book, BM25, dense, hybrid, oracle 증거 "
    "사다리를 비교했다. 검색 증거는 답변 정확도를 Qwen 0.3083에서 0.6650, "
    "Llama 0.2300에서 0.6800으로 높였고 oracle은 0.7467/0.6900이었다(표 13)."
)
CORRECTED_ANSWER = (
    "VRU 600문항의 고정 생성기에서 closed-book, distractor, dense, prefilter, "
    "oracle 증거 사다리를 비교했다. 검색 증거는 답변 정확도를 Qwen 0.3083에서 "
    "0.6650, Llama 0.3067에서 0.6650으로 높였고 prefilter는 0.6800/0.6667, "
    "oracle은 0.7467/0.6900이었다(표 13)."
)


def parse_xml(data: bytes) -> etree._Element:
    return etree.fromstring(data, etree.XMLParser(remove_blank_text=False))


def serialize_xml(root: etree._Element) -> bytes:
    return etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )


def first_child(parent: etree._Element, tag: str) -> etree._Element | None:
    return parent.find(tag)


def ensure_first(parent: etree._Element, tag: str) -> etree._Element:
    child = first_child(parent, tag)
    if child is None:
        child = etree.Element(tag)
        parent.insert(0, child)
    return child


def remove_theme_attrs(rfonts: etree._Element) -> None:
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(W + name, None)


def set_run_properties(rpr: etree._Element) -> None:
    rfonts = ensure_first(rpr, W + "rFonts")
    remove_theme_attrs(rfonts)
    rfonts.set(W + "ascii", LATIN_FONT)
    rfonts.set(W + "hAnsi", LATIN_FONT)
    rfonts.set(W + "eastAsia", KOREAN_FONT)
    rfonts.set(W + "cs", LATIN_FONT)
    rfonts.set(W + "hint", "eastAsia")

    lang = rpr.find(W + "lang")
    if lang is None:
        lang = etree.SubElement(rpr, W + "lang")
    lang.set(W + "val", "en-US")
    lang.set(W + "eastAsia", "ko-KR")
    lang.set(W + "bidi", "en-US")


def ensure_style_defaults(root: etree._Element) -> None:
    if root.tag != W + "styles":
        return
    defaults = root.find(W + "docDefaults")
    if defaults is None:
        defaults = etree.Element(W + "docDefaults")
        root.insert(0, defaults)
    rpr_default = defaults.find(W + "rPrDefault")
    if rpr_default is None:
        rpr_default = etree.SubElement(defaults, W + "rPrDefault")
    rpr = rpr_default.find(W + "rPr")
    if rpr is None:
        rpr = etree.SubElement(rpr_default, W + "rPr")
    set_run_properties(rpr)

    for style in root.findall(W + "style"):
        rpr = style.find(W + "rPr")
        if rpr is None:
            rpr = etree.SubElement(style, W + "rPr")
        set_run_properties(rpr)


def ensure_numbering_defaults(root: etree._Element) -> None:
    if root.tag != W + "numbering":
        return
    for lvl in root.xpath(".//w:lvl", namespaces=NS):
        rpr = lvl.find(W + "rPr")
        if rpr is None:
            rpr = etree.SubElement(lvl, W + "rPr")
        set_run_properties(rpr)


def disable_grid_snapping(root: etree._Element) -> None:
    for doc_grid in list(root.xpath(".//w:docGrid", namespaces=NS)):
        parent = doc_grid.getparent()
        if parent is not None:
            parent.remove(doc_grid)
    for ppr in root.xpath(".//w:pPr", namespaces=NS):
        snap = ppr.find(W + "snapToGrid")
        if snap is None:
            snap = etree.Element(W + "snapToGrid")
        else:
            ppr.remove(snap)
        snap.set(W + "val", "0")
        # In CT_PPr, snapToGrid must precede spacing/ind/jc/sectPr.  Appending
        # it after sectPr creates an invalid section-break paragraph that Word
        # may repair by discarding or reinterpreting the column definition.
        insertion_point = len(ppr)
        for tag in ("spacing", "ind", "contextualSpacing", "jc", "sectPr"):
            candidate = ppr.find(W + tag)
            if candidate is not None:
                insertion_point = min(insertion_point, ppr.index(candidate))
        ppr.insert(insertion_point, snap)


SECTPR_ORDER = {
    name: index
    for index, name in enumerate(
        (
            "headerReference",
            "footerReference",
            "footnotePr",
            "endnotePr",
            "type",
            "pgSz",
            "pgMar",
            "paperSrc",
            "pgBorders",
            "lnNumType",
            "pgNumType",
            "cols",
            "formProt",
            "vAlign",
            "noEndnote",
            "titlePg",
            "textDirection",
            "bidi",
            "rtlGutter",
            "docGrid",
            "printerSettings",
            "sectPrChange",
        )
    )
}


def normalize_section_properties(root: etree._Element) -> None:
    """Make continuous one/two-column sections deterministic in Word.

    The source layout uses temporary one-column sections so wide figures and
    tables can span the page.  Word is stricter than LibreOffice about the
    order of CT_SectPr children.  Explicit widths also prevent a different
    printer metric from recalculating the two columns.
    """

    for sectpr in root.xpath(".//w:sectPr", namespaces=NS):
        pg_sz = sectpr.find(W + "pgSz")
        pg_mar = sectpr.find(W + "pgMar")
        cols = sectpr.find(W + "cols")
        if cols is not None:
            page_width = int(
                pg_sz.get(W + "w", "11906") if pg_sz is not None else "11906"
            )
            left = int(
                pg_mar.get(W + "left", "850") if pg_mar is not None else "850"
            )
            right = int(
                pg_mar.get(W + "right", "850") if pg_mar is not None else "850"
            )
            text_width = page_width - left - right
            count = int(cols.get(W + "num", "1"))
            for child in list(cols):
                if child.tag == W + "col":
                    cols.remove(child)

            if count == 2:
                gap = int(cols.get(W + "space", "397"))
                first_width = (text_width - gap) // 2
                second_width = text_width - gap - first_width
                cols.set(W + "num", "2")
                cols.set(W + "equalWidth", "0")
                first = etree.SubElement(cols, W + "col")
                first.set(W + "w", str(first_width))
                first.set(W + "space", str(gap))
                second = etree.SubElement(cols, W + "col")
                second.set(W + "w", str(second_width))
                second.set(W + "space", "0")
            else:
                cols.set(W + "num", "1")
                cols.set(W + "equalWidth", "1")

        # CT_SectPr requires header/footer references before w:type.  The
        # generator's reversed order was tolerated by LibreOffice but is not
        # safe for Microsoft Word.
        children = list(sectpr)
        original_position = {id(child): pos for pos, child in enumerate(children)}
        children.sort(
            key=lambda child: (
                SECTPR_ORDER.get(etree.QName(child).localname, 10_000),
                original_position[id(child)],
            )
        )
        for child in children:
            sectpr.remove(child)
        for child in children:
            sectpr.append(child)


def replace_stale_answer(root: etree._Element) -> int:
    replacements = 0
    for text_node in root.xpath(".//w:t", namespaces=NS):
        if text_node.text and STALE_ANSWER in text_node.text:
            text_node.text = text_node.text.replace(STALE_ANSWER, CORRECTED_ANSWER)
            replacements += 1
    return replacements


def normalize_word_part(data: bytes) -> tuple[bytes, int]:
    root = parse_xml(data)
    ensure_style_defaults(root)
    ensure_numbering_defaults(root)

    for run in root.xpath(".//w:r", namespaces=NS):
        rpr = run.find(W + "rPr")
        if rpr is None:
            rpr = etree.Element(W + "rPr")
            run.insert(0, rpr)
        set_run_properties(rpr)

    for rpr in root.xpath(".//w:rPr", namespaces=NS):
        set_run_properties(rpr)

    disable_grid_snapping(root)
    normalize_section_properties(root)
    replacements = replace_stale_answer(root)
    return serialize_xml(root), replacements


def normalize_settings(data: bytes) -> bytes:
    root = parse_xml(data)
    for tag in ("embedTrueTypeFonts", "embedSystemFonts"):
        node = root.find(W + tag)
        if node is None:
            node = etree.Element(W + tag)
            root.insert(1, node)
        node.attrib.pop(W + "val", None)

    subset = root.find(W + "saveSubset")
    if subset is None:
        subset = etree.Element(W + "saveSubset")
        root.insert(2, subset)
    subset.set(W + "val", "0")

    spacing = root.find(W + "characterSpacingControl")
    if spacing is None:
        spacing = etree.SubElement(root, W + "characterSpacingControl")
    spacing.set(W + "val", "doNotCompress")

    theme_lang = root.find(W + "themeFontLang")
    if theme_lang is None:
        theme_lang = etree.SubElement(root, W + "themeFontLang")
    theme_lang.set(W + "val", "en-US")
    theme_lang.set(W + "eastAsia", "ko-KR")
    theme_lang.set(W + "bidi", "en-US")

    compat = root.find(W + "compat")
    if compat is None:
        compat = etree.SubElement(root, W + "compat")
    for child in list(compat):
        if child.tag == W + "compatSetting" and child.get(W + "name") == "compatibilityMode":
            compat.remove(child)
    mode = etree.SubElement(compat, W + "compatSetting")
    mode.set(W + "name", "compatibilityMode")
    mode.set(W + "uri", "http://schemas.microsoft.com/office/word")
    mode.set(W + "val", "15")
    return serialize_xml(root)


def normalize_theme(data: bytes) -> bytes:
    root = parse_xml(data)
    for branch in root.xpath(".//a:majorFont | .//a:minorFont", namespaces=NS):
        latin = branch.find(A + "latin")
        if latin is not None:
            latin.set("typeface", LATIN_FONT)
        ea = branch.find(A + "ea")
        if ea is not None:
            ea.set("typeface", KOREAN_FONT)
        cs = branch.find(A + "cs")
        if cs is not None:
            cs.set("typeface", LATIN_FONT)
        hang = None
        for font in branch.findall(A + "font"):
            if font.get("script") == "Hang":
                hang = font
                break
        if hang is None:
            hang = etree.SubElement(branch, A + "font")
            hang.set("script", "Hang")
        hang.set("typeface", KOREAN_FONT)
    return serialize_xml(root)


def check_embedding_permission(path: Path) -> None:
    font = TTFont(path)
    fs_type = int(font["OS/2"].fsType)
    if fs_type & 0x0002:
        raise RuntimeError(f"font forbids embedding: {path} fsType={fs_type}")
    if not (fs_type == 0 or fs_type & 0x0008):
        raise RuntimeError(
            f"font does not grant editable embedding: {path} fsType={fs_type}"
        )


def obfuscate_font(path: Path, key: uuid.UUID) -> bytes:
    data = bytearray(path.read_bytes())
    if len(data) < 32:
        raise RuntimeError(f"font file too small: {path}")
    key_bytes = key.bytes
    for idx in range(32):
        data[idx] ^= key_bytes[15 - (idx % 16)]
    return bytes(data)


def next_relationship_id(root: etree._Element) -> str:
    used = set()
    for rel in root.findall(REL + "Relationship"):
        rid = rel.get("Id", "")
        match = re.fullmatch(r"rId(\d+)", rid)
        if match:
            used.add(int(match.group(1)))
    candidate = 1
    while candidate in used:
        candidate += 1
    return f"rId{candidate}"


def font_table_with_embeddings(
    data: bytes,
    regular_rid: str,
    bold_rid: str,
    regular_key: uuid.UUID,
    bold_key: uuid.UUID,
) -> bytes:
    root = parse_xml(data)
    for font in list(root.findall(W + "font")):
        if font.get(W + "name") in {KOREAN_FONT, LATIN_FONT}:
            root.remove(font)

    latin = etree.SubElement(root, W + "font")
    latin.set(W + "name", LATIN_FONT)
    etree.SubElement(latin, W + "family").set(W + "val", "roman")
    etree.SubElement(latin, W + "pitch").set(W + "val", "variable")

    korean = etree.SubElement(root, W + "font")
    korean.set(W + "name", KOREAN_FONT)
    etree.SubElement(korean, W + "family").set(W + "val", "roman")
    etree.SubElement(korean, W + "pitch").set(W + "val", "variable")
    etree.SubElement(korean, W + "charset").set(W + "val", "81")

    regular = etree.SubElement(korean, W + "embedRegular")
    regular.set(R + "id", regular_rid)
    regular.set(W + "fontKey", "{" + str(regular_key).upper() + "}")
    regular.set(W + "subsetted", "0")

    bold = etree.SubElement(korean, W + "embedBold")
    bold.set(R + "id", bold_rid)
    bold.set(W + "fontKey", "{" + str(bold_key).upper() + "}")
    bold.set(W + "subsetted", "0")
    return serialize_xml(root)


def font_relationships(existing: bytes | None) -> tuple[bytes, str, str]:
    if existing is None:
        root = etree.Element(
            REL + "Relationships", nsmap={None: REL_NS}
        )
    else:
        root = parse_xml(existing)

    for rel in list(root.findall(REL + "Relationship")):
        if rel.get("Type", "").endswith("/font"):
            root.remove(rel)

    regular_rid = next_relationship_id(root)
    regular = etree.SubElement(root, REL + "Relationship")
    regular.set("Id", regular_rid)
    regular.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
    )
    regular.set("Target", "fonts/NanumMyeongjo-Regular.odttf")

    bold_rid = next_relationship_id(root)
    bold = etree.SubElement(root, REL + "Relationship")
    bold.set("Id", bold_rid)
    bold.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
    )
    bold.set("Target", "fonts/NanumMyeongjo-Bold.odttf")
    return serialize_xml(root), regular_rid, bold_rid


def content_types_with_font(data: bytes) -> bytes:
    root = parse_xml(data)
    for default in root.findall(CT + "Default"):
        if default.get("Extension") == "odttf":
            return serialize_xml(root)
    default = etree.SubElement(root, CT + "Default")
    default.set("Extension", "odttf")
    default.set(
        "ContentType",
        "application/vnd.openxmlformats-officedocument.obfuscatedFont",
    )
    return serialize_xml(root)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build(source: Path, output: Path, report: Path) -> None:
    for path in (source, KOREAN_REGULAR, KOREAN_BOLD):
        if not path.is_file():
            raise FileNotFoundError(path)
    check_embedding_permission(KOREAN_REGULAR)
    check_embedding_permission(KOREAN_BOLD)

    with zipfile.ZipFile(source, "r") as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    replacements = 0
    for name, data in list(parts.items()):
        if XML_PART_RE.match(name):
            parts[name], changed = normalize_word_part(data)
            replacements += changed

    parts["word/settings.xml"] = normalize_settings(parts["word/settings.xml"])
    if "word/theme/theme1.xml" in parts:
        parts["word/theme/theme1.xml"] = normalize_theme(
            parts["word/theme/theme1.xml"]
        )

    rel_name = "word/_rels/fontTable.xml.rels"
    rel_xml, regular_rid, bold_rid = font_relationships(parts.get(rel_name))
    parts[rel_name] = rel_xml

    regular_key = uuid.uuid4()
    bold_key = uuid.uuid4()
    parts["word/fontTable.xml"] = font_table_with_embeddings(
        parts["word/fontTable.xml"],
        regular_rid,
        bold_rid,
        regular_key,
        bold_key,
    )
    parts["word/fonts/NanumMyeongjo-Regular.odttf"] = obfuscate_font(
        KOREAN_REGULAR, regular_key
    )
    parts["word/fonts/NanumMyeongjo-Bold.odttf"] = obfuscate_font(
        KOREAN_BOLD, bold_key
    )
    parts["[Content_Types].xml"] = content_types_with_font(
        parts["[Content_Types].xml"]
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(
        output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    validation = validate(source, output, replacements)
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(
        json.dumps(validation, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    if validation["status"] != "PASS":
        raise RuntimeError(f"validation failed; see {report}")


def validate(source: Path, output: Path, replacements: int) -> dict[str, object]:
    checks: dict[str, object] = {}
    with zipfile.ZipFile(output, "r") as archive:
        names = set(archive.namelist())
        checks["zip_test"] = archive.testzip() is None
        checks["embedded_regular"] = (
            "word/fonts/NanumMyeongjo-Regular.odttf" in names
        )
        checks["embedded_bold"] = "word/fonts/NanumMyeongjo-Bold.odttf" in names
        checks["font_table_relationships"] = (
            "word/_rels/fontTable.xml.rels" in names
        )

        invalid_fonts: list[dict[str, str]] = []
        run_count = 0
        for name in sorted(names):
            if not XML_PART_RE.match(name):
                continue
            root = parse_xml(archive.read(name))
            for rfonts in root.xpath(".//w:rFonts", namespaces=NS):
                run_count += 1
                actual = {
                    key: rfonts.get(W + key)
                    for key in ("ascii", "hAnsi", "eastAsia", "cs")
                }
                expected = {
                    "ascii": LATIN_FONT,
                    "hAnsi": LATIN_FONT,
                    "eastAsia": KOREAN_FONT,
                    "cs": LATIN_FONT,
                }
                if actual != expected:
                    invalid_fonts.append({"part": name, "actual": str(actual)})
        checks["font_property_count"] = run_count
        checks["invalid_font_properties"] = invalid_fonts

        document = parse_xml(archive.read("word/document.xml"))
        text = "".join(document.itertext())
        checks["stale_answer_absent"] = "Llama 0.2300" not in text
        checks["corrected_answer_present"] = "Llama 0.3067" in text
        checks["doc_grid_removed"] = not bool(
            document.xpath(".//w:docGrid", namespaces=NS)
        )
        checks["grid_snap_disabled"] = all(
            node.get(W + "val") == "0"
            for node in document.xpath(".//w:snapToGrid", namespaces=NS)
        )
        section_break_paragraphs = document.xpath(
            ".//w:pPr[w:sectPr]", namespaces=NS
        )
        checks["section_break_property_order_valid"] = all(
            list(ppr)[-1].tag == W + "sectPr"
            for ppr in section_break_paragraphs
        )

        section_orders_valid = True
        two_column_widths_valid = True
        two_column_count = 0
        for sectpr in document.xpath(".//w:sectPr", namespaces=NS):
            known_order = [
                SECTPR_ORDER[etree.QName(child).localname]
                for child in sectpr
                if etree.QName(child).localname in SECTPR_ORDER
            ]
            if known_order != sorted(known_order):
                section_orders_valid = False

            cols = sectpr.find(W + "cols")
            if cols is None or cols.get(W + "num", "1") != "2":
                continue
            two_column_count += 1
            col_nodes = cols.findall(W + "col")
            pg_sz = sectpr.find(W + "pgSz")
            pg_mar = sectpr.find(W + "pgMar")
            if len(col_nodes) != 2 or pg_sz is None or pg_mar is None:
                two_column_widths_valid = False
                continue
            text_width = (
                int(pg_sz.get(W + "w"))
                - int(pg_mar.get(W + "left"))
                - int(pg_mar.get(W + "right"))
            )
            declared_width = sum(int(node.get(W + "w")) for node in col_nodes)
            declared_gap = int(col_nodes[0].get(W + "space", "0"))
            if (
                cols.get(W + "equalWidth") != "0"
                or declared_width + declared_gap != text_width
            ):
                two_column_widths_valid = False
        checks["section_property_order_valid"] = section_orders_valid
        checks["explicit_two_column_widths_valid"] = two_column_widths_valid
        checks["two_column_section_count"] = two_column_count

    from docx import Document

    source_doc = Document(source)
    output_doc = Document(output)
    checks["paragraph_count_preserved"] = len(source_doc.paragraphs) == len(
        output_doc.paragraphs
    )
    checks["table_count_preserved"] = len(source_doc.tables) == len(
        output_doc.tables
    )
    checks["inline_shape_count_preserved"] = len(source_doc.inline_shapes) == len(
        output_doc.inline_shapes
    )
    checks["section_count_preserved"] = len(source_doc.sections) == len(
        output_doc.sections
    )
    checks["answer_replacements"] = replacements
    checks["source_sha256"] = sha256(source)
    checks["output_sha256"] = sha256(output)
    checks["source_bytes"] = source.stat().st_size
    checks["output_bytes"] = output.stat().st_size
    checks["korean_font"] = KOREAN_FONT
    checks["latin_font"] = LATIN_FONT
    checks["korean_regular_fsType"] = int(TTFont(KOREAN_REGULAR)["OS/2"].fsType)
    checks["korean_bold_fsType"] = int(TTFont(KOREAN_BOLD)["OS/2"].fsType)

    boolean_checks = [
        checks["zip_test"],
        checks["embedded_regular"],
        checks["embedded_bold"],
        checks["font_table_relationships"],
        not checks["invalid_font_properties"],
        checks["stale_answer_absent"],
        checks["corrected_answer_present"],
        checks["doc_grid_removed"],
        checks["grid_snap_disabled"],
        checks["section_break_property_order_valid"],
        checks["section_property_order_valid"],
        checks["explicit_two_column_widths_valid"],
        checks["two_column_section_count"] > 0,
        checks["paragraph_count_preserved"],
        checks["table_count_preserved"],
        checks["inline_shape_count_preserved"],
        checks["section_count_preserved"],
        replacements == 1,
    ]
    return {
        "status": "PASS" if all(boolean_checks) else "FAIL",
        "checks": checks,
        "rule_7": {
            "page_limit": 20,
            "korean_title_pt": 18,
            "english_title_pt": 12,
            "author_pt": 8,
            "abstract_heading_pt": 12,
            "abstract_body_pt": 9,
            "body_heading_pt": 12,
            "body_pt": 9,
            "font_family_rule": "not specified by DBR; unified by user request",
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--report", type=Path, default=REPORT)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve(), args.report.resolve())
    print(f"editable docx: {args.output.resolve()}")
    print(f"verification: {args.report.resolve()}")


if __name__ == "__main__":
    main()
