#!/usr/bin/env python3
"""Verify the DBR v6 review-submission DOCX/PDF and report hard failures.

This checker covers the requirements that can be tested mechanically:
A4/page count, abstract and keyword bounds, anonymous review copy, numbered
captions, object order, typography, margins, table padding, image resolution,
embedded PDF fonts, and unresolved author-information placeholders.
"""

from __future__ import annotations

import re
import subprocess
import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
M = ROOT / "2026_KIISE" / "manuscript"
STEM = "kiise_dbr_manuscript_v6_submission_revision"
MD = M / f"{STEM}.md"
DOCX = M / f"{STEM}.docx"
PDF = M / f"{STEM}.pdf"

KOR_TITLE = (
    "도시 감시 멀티모달 데이터베이스의 증거 계층 설계: "
    "비순환 평가 기반 저장·검색·색인과 시각 언어 모델 답변 전파"
)
ENG_TITLE = (
    "Evidence-Layer Design for Multimodal Urban-Surveillance Databases: "
    "Non-Circular Evaluation of Storage, Retrieval, Indexing, and "
    "Vision-Language-Model Answer Propagation"
)
EXPECTED_FIGURES = 3
AUTHOR_LINES = {
    "박천복, Eduardo Linares, 김승현, 서영균†",
    "Cheonbok Park, Eduardo Linares, Seunghyun Kim, Young-Kyoon Suh†",
}
EXPECTED_LATIN = "Liberation Serif"
EXPECTED_KOREAN = "NanumMyeongjo"


class Checks:
    def __init__(self) -> None:
        self.passed: list[str] = []
        self.failed: list[str] = []
        self.pending: list[str] = []

    def require(self, condition: bool, label: str, detail: str = "") -> None:
        message = f"{label}{': ' + detail if detail else ''}"
        (self.passed if condition else self.failed).append(message)

    def note_pending(self, label: str) -> None:
        self.pending.append(label)


def command(*args: str) -> str:
    return subprocess.run(
        args, check=True, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE
    ).stdout


def paragraph_text(node) -> str:
    return re.sub(r"\s+", " ", "".join(node.itertext())).strip()


def first_nonempty_sibling(node, direction: str):
    cursor = getattr(node, direction)()
    while cursor is not None:
        if cursor.tag == qn("w:tbl"):
            return cursor
        if cursor.tag == qn("w:p"):
            if paragraph_text(cursor) or cursor.xpath(".//w:drawing"):
                return cursor
        cursor = getattr(cursor, direction)()
    return None


def run_size_pt(run) -> float | None:
    return None if run.font.size is None else round(run.font.size.pt, 2)


def all_text_run_sizes(paragraph) -> set[float | None]:
    return {run_size_pt(run) for run in paragraph.runs if run.text.strip()}


def main() -> int:
    checks = Checks()
    for path in (MD, DOCX, PDF):
        checks.require(path.is_file() and path.stat().st_size > 0, "산출물 존재", path.name)
    if checks.failed:
        return finish(checks)

    source = MD.read_text(encoding="utf-8")
    doc = Document(DOCX)

    # Review-copy metadata and abstract contract.
    page_parts = source.split("[DBR_PAGE_BREAK]", 1)
    checks.require(len(page_parts) == 2, "첫 페이지 뒤 명시적 페이지 나눔")
    if len(page_parts) == 2:
        checks.require(
            all(name in page_parts[0] for name in AUTHOR_LINES),
            "첫 페이지 저자 정보",
        )
        anonymized_front = page_parts[1].split("[DBR_BODY_TWO_COLUMNS]", 1)[0]
        checks.require(
            all(name not in anonymized_front for name in AUTHOR_LINES),
            "심사용 다음 페이지 익명화",
        )

    kor_match = re.search(r"## 초록\s+(.*?)\s+주요어:", source, flags=re.S)
    eng_match = re.search(r"## Abstract\s+(.*?)\s+Keywords:", source, flags=re.S)
    kor_chars = len(re.sub(r"\s+", "", kor_match.group(1))) if kor_match else 0
    eng_words = len(re.findall(r"\b[\w'-]+\b", eng_match.group(1))) if eng_match else 0
    checks.require(300 <= kor_chars <= 500, "한글 초록 300–500자", str(kor_chars))
    checks.require(100 <= eng_words <= 200, "영문 초록 100–200단어", str(eng_words))

    kor_kw = re.search(r"주요어:\s*([^\n]+)", source)
    eng_kw = re.search(r"Keywords:\s*([^\n]+)", source)
    kor_kw_count = len(kor_kw.group(1).split(",")) if kor_kw else 0
    eng_kw_count = len(eng_kw.group(1).split(",")) if eng_kw else 0
    checks.require(3 <= kor_kw_count <= 6, "한글 주요어 3–6개", str(kor_kw_count))
    checks.require(3 <= eng_kw_count <= 6, "영문 주요어 3–6개", str(eng_kw_count))
    checks.require(
        re.search(r"감사의\s*글|acknowledg(?:e)?ments?", source, flags=re.I) is None,
        "심사용 원고 감사의 글 미포함",
    )

    table_numbers = [
        int(value) for value in re.findall(r"\*\*\\?<표\s+(\d+)\\?>", source)
    ]
    figure_numbers = [
        int(value)
        for value in re.findall(r"(?m)^\*\*\\?<그림\s+(\d+)\\?>", source)
    ]
    checks.require(table_numbers == list(range(1, 13)), "표 번호 1–12 연속")
    checks.require(
        figure_numbers == list(range(1, EXPECTED_FIGURES + 1)),
        f"그림 번호 1–{EXPECTED_FIGURES} 연속",
    )
    checks.require(len(doc.tables) == 13, "DOCX 표 수", f"{len(doc.tables)} (저자표+결과표 12)")
    checks.require(
        len(doc.inline_shapes) == EXPECTED_FIGURES,
        "DOCX 그림 수",
        str(len(doc.inline_shapes)),
    )

    # A4, review page limit, and margins.
    pdfinfo = command("pdfinfo", str(PDF))
    page_match = re.search(r"^Pages:\s+(\d+)", pdfinfo, flags=re.M)
    size_match = re.search(
        r"^Page size:\s+([\d.]+) x ([\d.]+) pts \(([^)]+)\)", pdfinfo, flags=re.M
    )
    pages = int(page_match.group(1)) if page_match else 0
    checks.require(1 <= pages <= 20, "A4 심사용 20쪽 이내", str(pages))
    checks.require(
        bool(size_match)
        and abs(float(size_match.group(1)) - 595.3) < 1
        and abs(float(size_match.group(2)) - 841.9) < 1,
        "PDF A4 용지",
        size_match.group(0) if size_match else "미검출",
    )
    for index, section in enumerate(doc.sections, start=1):
        checks.require(
            abs(section.page_width.cm - 21.0) < 0.03
            and abs(section.page_height.cm - 29.7) < 0.03,
            f"DOCX section {index} A4",
        )
        checks.require(
            all(
                abs(value.cm - 1.5) < 0.03
                for value in (
                    section.top_margin,
                    section.bottom_margin,
                    section.left_margin,
                    section.right_margin,
                )
            ),
            f"DOCX section {index} 사방 여백 1.5 cm",
        )

    # Font sizes and object spacing/order.
    size_errors: list[str] = []
    font_errors: list[str] = []
    marker_errors: list[str] = []
    for paragraph in doc.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        if not text:
            continue
        sizes = all_text_run_sizes(paragraph)
        expected: float | None = None
        if text == KOR_TITLE:
            expected = 18
        elif text == ENG_TITLE or text in {"초록", "Abstract"}:
            expected = 12
        elif text in AUTHOR_LINES:
            expected = 8
        elif paragraph.style and paragraph.style.name.startswith("Heading"):
            expected = 12
        elif text.startswith(("<표 ", "<그림 ")) or re.match(r"^\[\d+\]", text):
            expected = 8
        else:
            expected = 9
        if sizes != {expected}:
            size_errors.append(f"{text[:28]}={sorted(str(v) for v in sizes)}≠{expected}")
        if text.startswith("[DBR_"):
            marker_errors.append(text)
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            r_fonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
            if r_fonts is None:
                font_errors.append(f"{run.text[:20]}: rFonts 없음")
                continue
            latin = r_fonts.get(qn("w:ascii"))
            hansi = r_fonts.get(qn("w:hAnsi"))
            east_asia = r_fonts.get(qn("w:eastAsia"))
            if latin != EXPECTED_LATIN or hansi != EXPECTED_LATIN or east_asia != EXPECTED_KOREAN:
                font_errors.append(
                    f"{run.text[:20]}: ascii={latin}, hAnsi={hansi}, eastAsia={east_asia}"
                )
    checks.require(not size_errors, "본문·제목·캡션 지정 글자 크기", "; ".join(size_errors[:3]))
    checks.require(not font_errors, "DOCX 본문 지정 글꼴", "; ".join(font_errors[:3]))
    checks.require(not marker_errors, "DOCX 편집 마커 제거", "; ".join(marker_errors))

    body = doc._element.body
    table_order_errors: list[str] = []
    figure_order_errors: list[str] = []
    result_table_index = 0
    for child in body.iterchildren():
        if child.tag == qn("w:tbl"):
            result_table_index += 1
            previous = first_nonempty_sibling(child, "getprevious")
            previous_text = paragraph_text(previous) if previous is not None else ""
            # The author-information table intentionally has no numbered caption.
            if result_table_index > 1 and not previous_text.startswith("<표 "):
                table_order_errors.append(previous_text[:50] or "(없음)")
        elif child.tag == qn("w:p") and child.xpath(".//w:drawing"):
            following = first_nonempty_sibling(child, "getnext")
            following_text = paragraph_text(following) if following is not None else ""
            if not following_text.startswith("<그림 "):
                figure_order_errors.append(following_text[:50] or "(없음)")
    checks.require(not table_order_errors, "표 제목이 표 위에 배치", "; ".join(table_order_errors))
    checks.require(not figure_order_errors, "그림 캡션이 그림 아래 배치", "; ".join(figure_order_errors))

    caption_spacing_errors: list[str] = []
    figure_spacing_errors: list[str] = []
    for paragraph in doc.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text).strip()
        fmt = paragraph.paragraph_format
        before = round(fmt.space_before.pt, 2) if fmt.space_before else 0
        after = round(fmt.space_after.pt, 2) if fmt.space_after else 0
        if text.startswith("<표 ") and (before != 8 or after != 4):
            caption_spacing_errors.append(f"{text[:12]}={before}/{after}pt")
        if text.startswith("<그림 ") and (before != 1.5 or after != 7):
            caption_spacing_errors.append(f"{text[:12]}={before}/{after}pt")
        if paragraph._p.xpath(".//w:drawing") and (before != 8 or after != 1.5):
            figure_spacing_errors.append(f"{before}/{after}pt")
    checks.require(
        not caption_spacing_errors,
        "표·그림 캡션 여백",
        "; ".join(caption_spacing_errors),
    )
    checks.require(
        not figure_spacing_errors,
        "그림과 캡션 사이 여백",
        "; ".join(figure_spacing_errors),
    )

    # Table cell padding and fixed row integrity.
    table_margin_errors: list[str] = []
    for table_index, table in enumerate(doc.tables):
        for row in table.rows:
            if not row._tr.xpath("./w:trPr/w:cantSplit"):
                table_margin_errors.append(f"table {table_index}: 행 분할 방지 없음")
            for cell in row.cells:
                for side, expected in (("top", "85"), ("bottom", "85"), ("left", "70"), ("right", "70")):
                    nodes = cell._tc.xpath(f"./w:tcPr/w:tcMar/w:{side}")
                    if not nodes or nodes[0].get(qn("w:w")) != expected:
                        table_margin_errors.append(
                            f"table {table_index}: {side}={nodes[0].get(qn('w:w')) if nodes else '없음'}"
                        )
    checks.require(
        not table_margin_errors,
        "표 셀 여백·행 분할 방지",
        "; ".join(table_margin_errors[:3]),
    )

    # Figure native resolution and PDF export resolution.
    dpi_values: list[float] = []
    width_values: list[float] = []
    for shape in doc.inline_shapes:
        blips = shape._inline.xpath(".//a:blip")
        if len(blips) != 1:
            continue
        rel_id = blips[0].get(qn("r:embed"))
        blob = doc.part.related_parts[rel_id].blob
        with Image.open(BytesIO(blob)) as image:
            dpi_x = image.width / (shape.width / 914400)
            dpi_y = image.height / (shape.height / 914400)
            dpi_values.append(min(dpi_x, dpi_y))
            width_values.append(shape.width / 360000)
    checks.require(
        len(dpi_values) == EXPECTED_FIGURES and min(dpi_values) >= 300,
        "DOCX 그림 유효 해상도 ≥300 dpi",
        f"최소 {min(dpi_values):.1f}" if dpi_values else "미검출",
    )
    checks.require(
        len(width_values) == EXPECTED_FIGURES and max(width_values) <= 15.8 + 0.03,
        "그림 폭이 전폭 본문 영역 이내",
        f"최대 {max(width_values):.2f} cm" if width_values else "미검출",
    )

    image_list = command("pdfimages", "-list", str(PDF))
    pdf_rasters = [
        line
        for line in image_list.splitlines()
        if re.match(r"\s*\d+\s+\d+\s+image\s+", line)
    ]
    pdf_ppi = []
    for line in pdf_rasters:
        fields = line.split()
        pdf_ppi.append(min(int(fields[12]), int(fields[13])))
    checks.require(
        len(pdf_rasters) == EXPECTED_FIGURES and min(pdf_ppi) >= 300,
        f"PDF 그림 {EXPECTED_FIGURES}개·출력 해상도 ≥300 ppi",
        f"최소 {min(pdf_ppi)}" if pdf_ppi else "미검출",
    )

    fonts = command("pdffonts", str(PDF)).splitlines()[2:]
    font_rows = [line.split() for line in fonts if line.strip()]
    embedded = bool(font_rows) and all(len(row) >= 6 and row[4] == "yes" for row in font_rows)
    checks.require(embedded, "PDF 사용 글꼴 전부 embedded")
    checks.require(
        any("NanumMyeongjo" in row[0] for row in font_rows)
        and any("LiberationSerif" in row[0] for row in font_rows),
        "PDF 한글·라틴 본문 글꼴 포함",
    )

    # These are administrative submission blockers, not layout failures.
    if "[직위 최종 확인 필요]" in source:
        checks.note_pending("박천복·Eduardo Linares·김승현의 직위")
    if "[전화번호 최종 입력 필요]" in source:
        checks.note_pending("교신저자 전화번호")

    return finish(checks)


def finish(checks: Checks) -> int:
    for message in checks.passed:
        print(f"PASS  {message}")
    for message in checks.failed:
        print(f"FAIL  {message}")
    for message in checks.pending:
        print(f"PENDING  {message}")
    status = "PASS_WITH_METADATA_PENDING" if not checks.failed and checks.pending else (
        "PASS" if not checks.failed else "FAIL"
    )
    print(
        f"SUMMARY  status={status} pass={len(checks.passed)} "
        f"fail={len(checks.failed)} pending={len(checks.pending)}"
    )
    return 1 if checks.failed else 0


if __name__ == "__main__":
    sys.exit(main())
